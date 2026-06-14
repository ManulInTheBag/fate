"""Write the aggregated stats to a multi-sheet .xlsx workbook."""

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import config
import stats

HEADER_FILL = PatternFill("solid", fgColor="1F1F1F")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb, title, columns, rows):
    ws = wb.create_sheet(title)
    headers = [c[1] for c in columns]
    keys = [c[0] for c in columns]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([r.get(k, "") for k in keys])
    # widths + freeze header
    for i, h in enumerate(headers, 1):
        width = max(len(str(h)), *(len(str(r.get(keys[i - 1], ""))) for r in rows)) if rows else len(h)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 40)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def export(path=None, filters=None, period_label="All time"):
    path = path or config.EXPORT_PATH
    filters = filters or {}
    data = stats.compute(**filters)

    wb = Workbook()
    wb.remove(wb.active)  # drop default sheet

    # Overview
    ov = wb.create_sheet("Overview")
    ov.append(["Fate Scoreboard Stats"])
    ov["A1"].font = Font(bold=True, size=14)
    ov.append(["Period", period_label])
    ov.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ov.append(["Matches", len(data["matches"])])
    ov.append(["Players tracked", len(data["players"])])
    ov.append(["Heroes seen", len(data["heroes"])])
    ov.column_dimensions["A"].width = 18
    ov.column_dimensions["B"].width = 24

    _sheet(wb, "Player Stats", [
        ("player", "Player"), ("games", "Games"), ("wins", "Wins"), ("losses", "Losses"),
        ("winrate", "Win %"), ("kills", "Kills"), ("deaths", "Deaths"), ("assists", "Assists"),
        ("kda", "KDA"), ("avg_k", "Avg K"), ("avg_d", "Avg D"), ("avg_a", "Avg A"),
        ("heroes_played", "Heroes"), ("most_played_hero", "Most Played"),
    ], data["players"])

    _sheet(wb, "Hero Stats", [
        ("hero", "Hero"), ("picks", "Picks"), ("wins", "Wins"), ("losses", "Losses"),
        ("winrate", "Win %"), ("avg_k", "Avg K"), ("avg_d", "Avg D"), ("avg_a", "Avg A"),
        ("kda", "KDA"),
    ], data["heroes"])

    _sheet(wb, "Player-Hero", [
        ("player", "Player"), ("hero", "Hero"), ("games", "Games"), ("wins", "Wins"),
        ("winrate", "Win %"), ("kda", "KDA"),
    ], data["player_hero"])

    _sheet(wb, "Teammates", [
        ("player_a", "Player A"), ("player_b", "Player B"), ("games_together", "Games"),
        ("wins_together", "Wins"), ("winrate", "Win %"),
    ], data["teammates"])

    matches = [{
        **m,
        "date": datetime.fromtimestamp(m["posted_at"]).strftime("%Y-%m-%d %H:%M"),
        "score": f'{m["score_top"]}-{m["score_bottom"]}',
    } for m in data["matches"]]
    _sheet(wb, "Matches", [
        ("match_id", "ID"), ("date", "Date"), ("winner", "Winner"), ("score", "Score"),
    ], matches)

    raw = [{
        **r,
        "date": datetime.fromtimestamp(r["posted_at"]).strftime("%Y-%m-%d %H:%M"),
        "result": "Win" if r["won"] else "Loss",
    } for r in data["raw"]]
    _sheet(wb, "Raw Match Players", [
        ("match_id", "Match"), ("date", "Date"), ("team", "Team"), ("result", "Result"),
        ("player", "Player"), ("hero", "Hero"), ("kills", "K"), ("deaths", "D"),
        ("assists", "A"),
    ], raw)

    wb.save(path)
    return path, len(data["matches"])


if __name__ == "__main__":
    import db
    db.init_db()
    print("Wrote", export())
